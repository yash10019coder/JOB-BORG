"""Reading an uploaded document safely, and normalizing its text (rules D1-D11, N1-N7).

``read_document`` takes the raw bytes and the declared extension, checks the file
for what it really is, rejects anything risky, and returns plain text. It never
OCRs, never executes anything, and never lets exception text escape: callers get
a :class:`DocumentError` carrying a short ``code`` (D11), because exception
messages from parsers can contain resume content.

Pure and dependency-light (pypdf, python-docx, stdlib): no database, no network.
"""
import io
import re
import unicodedata
import zipfile
from collections import Counter
from dataclasses import dataclass

from django.conf import settings

ALLOWED_EXTENSIONS = (".pdf", ".docx", ".txt")
MIN_TEXT_CHARS = 200
DOCX_MAX_ENTRIES = 200
DOCX_MAX_UNCOMPRESSED = 20 * 1024 * 1024
DOCX_MAX_RATIO = 100
DOCX_RATIO_FLOOR = 512 * 1024  # tiny, highly repetitive XML parts are not bombs

BULLET_GLYPHS = "•●▪■◦○∙·‣⁃▸►"
_ACTIVE_ACTIONS = frozenset({"/JavaScript", "/Launch", "/SubmitForm", "/ImportData"})
_RAW_ACTIVE = re.compile(rb"/(?:JavaScript|JS|Launch|EmbeddedFiles|RichMedia|XFA)(?![A-Za-z])")


class DocumentError(Exception):
    """The document cannot be imported. ``code`` is short and safe to show/log."""

    def __init__(self, code):
        super().__init__(code)
        self.code = code


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    kind: str
    pages: int


def _setting(name, default):
    return getattr(settings, name, default)


# --------------------------------------------------------------------------
# Intake (D rules)
# --------------------------------------------------------------------------
def read_document(data, extension):
    """Return the document's plain text, or raise :class:`DocumentError`."""
    extension = (extension or "").lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise DocumentError("unsupported_type")  # D1
    if not data:
        raise DocumentError("empty_file")  # D2
    if len(data) > _setting("PROFILE_IMPORT_MAX_PDF_BYTES", 5 * 1024 * 1024):
        raise DocumentError("file_too_large")  # D2

    if extension == ".pdf":
        text, pages = _read_pdf(data)
        kind = "pdf"
    elif extension == ".docx":
        text, pages = _read_docx(data), 1
        kind = "docx"
    else:
        text, pages = _read_txt(data), 1
        kind = "txt"

    if len(re.sub(r"\s", "", text)) < MIN_TEXT_CHARS:
        raise DocumentError("no_text_found")  # D6: scanned/image files; no OCR, ever
    return ExtractedDocument(text=text, kind=kind, pages=pages)


def _read_txt(data):
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise DocumentError("txt_unreadable") from None
    if "\x00" in text:
        raise DocumentError("txt_unreadable")
    return text


def _resolve(obj):
    return obj.get_object() if hasattr(obj, "get_object") else obj


def _action_is_active(action, depth=0):
    """True for an action that runs code or launches something (D4). A plain view
    destination (an array) or a URI/GoTo action is fine."""
    action = _resolve(action)
    if depth > 8 or not hasattr(action, "get") or isinstance(action, (list, str, bytes)):
        return False
    if "/JS" in action or action.get("/S") in _ACTIVE_ACTIONS:
        return True
    following = action.get("/Next")
    if following is None:
        return False
    following = _resolve(following)
    items = following if isinstance(following, list) else [following]
    return any(_action_is_active(item, depth + 1) for item in items)


def _get(mapping, key):
    """``mapping[key]`` with indirect references resolved, or ``None``."""
    value = mapping.get(key)
    return None if value is None else _resolve(value)


def _any_active(mapping):
    mapping = _resolve(mapping) if mapping is not None else None
    if mapping is None or not hasattr(mapping, "values"):
        return False
    return any(_action_is_active(value) for value in mapping.values())


def _pdf_has_active_content(reader):
    root = _resolve(reader.trailer["/Root"])
    # A bare /OpenAction view destination is normal (LaTeX/hyperref writes one):
    # only an action that runs code counts (D4).
    open_action = root.get("/OpenAction")
    if open_action is not None and _action_is_active(open_action):
        return True
    if _any_active(root.get("/AA")):
        return True
    names = _get(root, "/Names")
    if names is not None and ("/JavaScript" in names or "/EmbeddedFiles" in names):
        return True
    form = _get(root, "/AcroForm")
    if form is not None and "/XFA" in form:
        return True
    for page in reader.pages:
        if _any_active(page.get("/AA")):
            return True
        for annotation in _get(page, "/Annots") or []:
            annotation = _resolve(annotation)
            action = annotation.get("/A")
            if action is not None and _action_is_active(action):
                return True
            if _any_active(annotation.get("/AA")):
                return True
    return False


def _read_pdf(data):
    from pypdf import PdfReader

    if not data.startswith(b"%PDF-"):  # D1: sniffed, not trusted from the name
        raise DocumentError("not_a_pdf")
    if _RAW_ACTIVE.search(data):  # D4 backstop for tokens that are never benign
        raise DocumentError("pdf_active_content")
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DocumentError("pdf_encrypted")  # D3, even owner-password-only
        page_count = len(reader.pages)
        if page_count > _setting("PROFILE_IMPORT_MAX_PAGES", 10):
            raise DocumentError("pdf_too_long")  # D3
        if _pdf_has_active_content(reader):
            raise DocumentError("pdf_active_content")  # D4
        parts = []
        for page in reader.pages:
            try:
                parts.append(page.extract_text() or "")
            except Exception:  # noqa: BLE001 -- one bad page must not fail the file
                continue
    except DocumentError:
        raise
    except Exception:  # noqa: BLE001 -- parser errors can carry content (D11)
        raise DocumentError("pdf_unreadable") from None
    return "\n".join(parts), page_count


def _inspect_docx(data):
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
        infos = archive.infolist()
    except zipfile.BadZipFile:
        raise DocumentError("docx_invalid") from None
    if len(infos) > DOCX_MAX_ENTRIES:
        raise DocumentError("docx_too_complex")
    total = 0
    names = set()
    for info in infos:
        name = info.filename
        if name.startswith("/") or ".." in name.split("/"):
            raise DocumentError("docx_unsafe")
        total += info.file_size
        if total > DOCX_MAX_UNCOMPRESSED:
            raise DocumentError("docx_too_large")
        if (
            info.compress_size
            and info.file_size > DOCX_RATIO_FLOOR
            and info.file_size / info.compress_size > DOCX_MAX_RATIO
        ):
            raise DocumentError("docx_too_large")
        names.add(name)
    if any(name.lower().endswith("vbaproject.bin") for name in names):
        raise DocumentError("docx_macros")
    if "word/document.xml" not in names:
        raise DocumentError("docx_invalid")


def _read_docx(data):
    from docx import Document

    _inspect_docx(data)
    try:
        document = Document(io.BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append("\t".join(cell.text for cell in row.cells))
    except Exception:  # noqa: BLE001 -- parser errors can carry content (D11)
        raise DocumentError("docx_unreadable") from None
    return "\n".join(parts)


# --------------------------------------------------------------------------
# Normalization (N rules)
# --------------------------------------------------------------------------
_CID = re.compile(r"\(cid:\d+\)")
_INVISIBLE = re.compile(r"[​-‏⁠﻿­-]")
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]")


@dataclass(frozen=True)
class NormalizedText:
    text: str
    lines: tuple
    noise: frozenset  # indexes of repeated header/footer lines (N7)
    truncated: bool


def bullet_of(line):
    """The bullet marker a line starts with, or ``None`` (N4)."""
    stripped = line.lstrip()
    if stripped and stripped[0] in BULLET_GLYPHS:
        return stripped[0]
    if re.match(r"^[-*]\s+\S", stripped):
        return stripped[0]
    return None


def strip_bullet(line):
    """The line without its leading bullet marker (N4)."""
    stripped = line.lstrip()
    if stripped and stripped[0] in BULLET_GLYPHS:
        return stripped[1:].lstrip()
    if re.match(r"^[-*]\s+\S", stripped):
        return stripped[1:].lstrip()
    return stripped


def normalize_text(text):
    """NFKC, strip invisible/control junk, tidy whitespace, cap the length and
    mark repeated page headers/footers (N1-N7). Never re-flows lines (N3)."""
    text = unicodedata.normalize("NFKC", text or "")  # N1 (also NBSP -> space, ligatures)
    text = _CID.sub("", text)
    text = _INVISIBLE.sub("", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _CONTROL.sub("", text)
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]  # N2
    collapsed, blanks = [], 0
    for line in lines:
        blanks = blanks + 1 if not line else 0
        if blanks <= 1:  # N2: three or more newlines become two
            collapsed.append(line)
    cap = _setting("PROFILE_IMPORT_MAX_TEXT_CHARS", 40_000)
    truncated, kept, size = False, [], 0
    for line in collapsed:  # N5: cut at a line boundary
        if size + len(line) + 1 > cap:
            truncated = True
            break
        kept.append(line)
        size += len(line) + 1
    counts = Counter(line.casefold() for line in kept if len(line) >= 6)
    seen, noise = set(), set()
    for index, line in enumerate(kept):  # N7: repeats after the first occurrence
        key = line.casefold()
        if counts.get(key, 0) >= 3:
            if key in seen:
                noise.add(index)
            seen.add(key)
    return NormalizedText(
        text="\n".join(kept), lines=tuple(kept), noise=frozenset(noise), truncated=truncated
    )
